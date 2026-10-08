from pathlib import Path

import openpyxl
import pytest

from merge_excel.profile import Profile
from merge_excel.scanner import Scan, Template


@pytest.fixture
def source(tmp_path: Path):
    folder = tmp_path / "input"
    folder.mkdir()

    def add(name: str, sheets: dict[str, list[list]], merges: dict[str, list[str]] | None = None) -> Path:
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for title, rows in sheets.items():
            sheet = workbook.create_sheet(title)
            for row in rows:
                sheet.append(row)
            for area in (merges or {}).get(title, []):
                sheet.merge_cells(area)
        workbook.save(folder / name)
        return folder

    return add


def template_with(scan: Scan, name: str) -> Template:
    return next(template for template in scan.templates if name in template.names)


def link(profile: Profile, scan: Scan, template: Template, mapping: dict[str, str]) -> None:
    for key, column_name in mapping.items():
        target = next(column for column in profile.columns if column.name == column_name)
        profile.link(template, scan.templates, key, target.id)


def read(path: Path) -> list[list]:
    return [list(row) for row in openpyxl.load_workbook(path)["Result"].iter_rows(values_only=True)]
