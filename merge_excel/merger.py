import re
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .convert import Transform, to_datetime, transformer
from .profile import Profile, ResultColumn
from .reader import open_book
from .scanner import FILE_KEY, SHEET_KEY, Header, Progress, Scan, Sheet, Template, has_digit, is_skipped, text

CELL_LIMIT = 32767
WIDTH_RANGE = (8, 60)
DATE_FORMAT = "dd.mm.yyyy"
DATETIME_FORMAT = "dd.mm.yyyy hh:mm:ss"
UNSAFE_NAME = re.compile(r'[\\/:*?"<>|]+')
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="305496")
TOP = Alignment(vertical="top")
WRAP = Alignment(vertical="top", wrap_text=True)


@dataclass(frozen=True)
class Source:
    indexes: tuple[int, ...]
    constants: tuple[str, ...]


@dataclass
class Record:
    key: object
    values: list[list]


def _sources(keys: list[list[str]], sheet: Sheet) -> list[Source]:
    constants = {FILE_KEY: sheet.path.name, SHEET_KEY: sheet.name}
    return [
        Source(tuple(sorted(int(k) for k in group if k not in constants)), tuple(constants[k] for k in group if k in constants))
        for group in keys
    ]


def _records(
    rows: Iterable[list], header: Header, width: int, sources: list[Source], transforms: list[Transform], skip: tuple[str, ...]
) -> list[Record]:
    records: list[Record] = []
    for row in rows:
        cells = row[header.left:header.left + width]
        if all(value is None for value in cells) or is_skipped(cells[0], skip):
            continue
        if cells[0] is not None:
            constants = [[v for v in map(t, s.constants) if v is not None] for s, t in zip(sources, transforms)]
            records.append(Record(cells[0], constants))
        elif not records:
            continue
        for values, source, transform in zip(records[-1].values, sources, transforms):
            parts = [part for part in (transform(cells[i]) for i in source.indexes if i < len(cells)) if part is not None]
            if parts:
                values.append(parts[0] if len(parts) == 1 else " ".join(map(text, parts)))
    while header.numbered and records and not has_digit(records[-1].key):
        records.pop()
    return records


def _value(values: list, column: ResultColumn):
    if column.join:
        return column.separator.join(map(text, values)) or None
    return values[0] if values else None


def merge(scan: Scan, profile: Profile, progress: Progress) -> list[list]:
    columns = profile.columns
    position = {column.id: index for index, column in enumerate(columns)}
    jobs: dict[Path, dict[str, tuple[Template, Sheet, list[list[str]]]]] = {}
    for template in scan.templates:
        keys: list[list[str]] = [[] for _ in columns]
        for key, target in profile.links_for(template, scan.templates).items():
            keys[position[target]].append(key)
        if any(keys):
            for sheet in template.sheets:
                jobs.setdefault(sheet.path, {})[sheet.name] = (template, sheet, keys)
    transforms = [transformer(column) for column in columns]
    total = sum(map(len, jobs.values()))
    done = 0
    output: list[list] = []
    for path in scan.files:
        sheets = jobs.get(path)
        if not sheets:
            continue
        try:
            with open_book(path) as book:
                for name in book.sheet_names:
                    if name not in sheets:
                        continue
                    template, sheet, keys = sheets[name]
                    progress(done, total, f"{path.name} — {name}")
                    done += 1
                    records = _records(
                        book.rows(name, sheet.header.body),
                        sheet.header,
                        template.width,
                        _sources(keys, sheet),
                        transforms,
                        profile.skip_prefixes,
                    )
                    output.extend([_value(values, column) for values, column in zip(record.values, columns)] for record in records)
        except Exception as error:
            raise RuntimeError(f"Cannot read {path.name}: {error}") from error
    return output


def _cell(sheet, value, alignment: Alignment):
    if value is None:
        return None
    if isinstance(value, str) and len(value) > CELL_LIMIT:
        value = value[:CELL_LIMIT]
    cell = WriteOnlyCell(sheet, value)
    cell.alignment = alignment
    if isinstance(value, date):
        cell.number_format = DATETIME_FORMAT if isinstance(value, datetime) and value.time() != time.min else DATE_FORMAT
    return cell


def _header_cell(sheet, title: str) -> WriteOnlyCell:
    cell = WriteOnlyCell(sheet, title)
    cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    return cell


def _write(path: Path, columns: list[ResultColumn], rows: list[list], header: bool = True) -> None:
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Result")
    for index, column in enumerate(columns):
        values = ([column.name] if header else []) + [row[index] for row in rows]
        longest = max((len(line) for value in values for line in text(value).splitlines()), default=0)
        sheet.column_dimensions[get_column_letter(index + 1)].width = min(max(longest + 2, WIDTH_RANGE[0]), WIDTH_RANGE[1])
    if header:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
        sheet.append([_header_cell(sheet, column.name) for column in columns])
    alignments = [WRAP if column.join else TOP for column in columns]
    for row in rows:
        sheet.append([_cell(sheet, value, alignment) for value, alignment in zip(row, alignments)])
    book.save(path)


def _is_blank(value) -> bool:
    return not any(character.isalnum() for character in text(value))


def _unique(rows: list[list], index: int) -> list:
    found = {}
    for row in rows:
        if not _is_blank(row[index]):
            found.setdefault(text(row[index]), row[index])
    return [found[key] for key in sorted(found, key=str.casefold)]


def _latest(rows: list[list], key: int, by: int) -> list[list]:
    best: dict[str, tuple[datetime | None, list]] = {}
    for row in rows:
        if _is_blank(row[key]):
            continue
        when, known = to_datetime(row[by]), best.get(text(row[key]))
        if known is None or (when is not None and (known[0] is None or when > known[0])):
            best[text(row[key])] = (when, row)
    return [row for _, row in sorted(best.values(), key=lambda item: item[0] or datetime.min, reverse=True)]


def _sibling(path: Path, suffix: str) -> Path:
    return path.with_name(f"{path.stem} - {UNSAFE_NAME.sub('_', suffix)}{path.suffix}")


@dataclass(frozen=True)
class Output:
    label: str
    rows: int
    path: Path


def export(scan: Scan, profile: Profile, path: Path, progress: Progress) -> list[Output]:
    rows = merge(scan, profile, progress)
    progress(1, 1, "Saving files…")
    columns = profile.columns
    _write(path, columns, rows)
    outputs = [Output("Merged records", len(rows), path)]
    unique = profile.column(profile.unique)
    if unique:
        key = columns.index(unique)
        values = _unique(rows, key)
        target = _sibling(path, f"unique {unique.name}")
        _write(target, [unique], [[value] for value in values], header=False)
        outputs.append(Output(f"Unique {unique.name}", len(values), target))
        by = profile.column(profile.latest_by)
        if profile.latest and by:
            latest = _latest(rows, key, columns.index(by))
            target = _sibling(path, f"latest by {by.name}")
            _write(target, columns, latest)
            outputs.append(Output(f"Latest record per {unique.name}", len(latest), target))
    return outputs
